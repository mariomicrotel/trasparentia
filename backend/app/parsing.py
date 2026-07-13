"""Estrazione testo dai documenti: PDF nativi (pypdf) con fallback OCR
(Tesseract italiano) per scansioni e immagini. Import lazy e tutto in try/except
così l'assenza di una dipendenza non blocca l'acquisizione."""
import io

from .config import settings


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def _ocr_pdf(data: bytes) -> str:
    from pdf2image import convert_from_bytes
    import pytesseract
    images = convert_from_bytes(data, dpi=200)
    return "\n".join(pytesseract.image_to_string(im, lang="ita") for im in images[:10])


def _ocr_image(data: bytes) -> str:
    from PIL import Image
    import pytesseract
    return pytesseract.image_to_string(Image.open(io.BytesIO(data)), lang="ita")


def _docx_text(data: bytes) -> str:
    """Testo da .docx preservando i ritorni a capo tra paragrafi: fondamentale
    per il parser normativo, che riconosce gli articoli su righe separate."""
    from docx import Document
    doc = Document(io.BytesIO(data))
    parti = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            parti.append(" \t ".join(c.text for c in row.cells))
    return "\n".join(parti)


def extract_text(filename: str, content_type: str, data: bytes) -> dict:
    name = (filename or "").lower()
    ct = (content_type or "").lower()
    try:
        if name.endswith(".pdf") or ct == "application/pdf":
            text = _pdf_text(data)
            ocr = False
            if len(text.strip()) < 40 and settings.OCR_ENABLED:
                try:
                    t2 = _ocr_pdf(data)
                    if len(t2.strip()) > len(text.strip()):
                        text, ocr = t2, True
                except Exception:
                    pass
            return {"text": text.strip(), "ocr": ocr}
        if name.endswith(".docx") or ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return {"text": _docx_text(data).strip(), "ocr": False}
        if name.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp")) or ct.startswith("image/"):
            if settings.OCR_ENABLED:
                return {"text": _ocr_image(data).strip(), "ocr": True}
            return {"text": "", "ocr": False}
        if name.endswith((".xml", ".txt", ".eml", ".csv", ".json", ".html")) or ct.startswith("text"):
            return {"text": data.decode("utf-8", "ignore").strip(), "ocr": False}
        return {"text": "", "ocr": False}
    except Exception as e:
        return {"text": "", "ocr": False, "error": str(e)}
